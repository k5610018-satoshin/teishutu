#!/usr/bin/env python3
"""サンプルのExcelデータを生成するスクリプト。

ロイロノートの「提出状況ダウンロード」で取得するExcelファイルに似た
サンプルデータを生成します。動作確認やデモ用に使用してください。
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


SAMPLE_STUDENTS = [
    "青木 太郎", "石田 花子", "上田 一郎", "遠藤 美咲", "大野 健太",
    "加藤 さくら", "木村 翔太", "黒田 優子", "小林 大輝", "佐藤 葵",
    "鈴木 蓮", "高橋 陽菜", "田中 悠人", "中村 結衣", "野田 颯太",
    "橋本 凛", "藤井 朝陽", "松本 七海", "山田 壮真", "渡辺 ひなた",
]


def create_sample_excel(output_dir: str = "./downloads"):
    """サンプルExcelファイルを生成する。"""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    # --- 数学のExcel ---
    _create_subject_file(
        filepath=out / "1年A組_数学_提出状況.xlsx",
        subject="数学",
        boxes={
            "Unit1 確認テスト": {
                "submitted": SAMPLE_STUDENTS[:15],
                "unsubmitted": SAMPLE_STUDENTS[15:],
            },
            "Unit2 演習プリント": {
                "submitted": SAMPLE_STUDENTS[:12],
                "unsubmitted": SAMPLE_STUDENTS[12:],
            },
            "Unit3 宿題": {
                "submitted": SAMPLE_STUDENTS[:18],
                "unsubmitted": SAMPLE_STUDENTS[18:],
            },
        },
    )

    # --- 英語のExcel ---
    _create_subject_file(
        filepath=out / "1年A組_英語_提出状況.xlsx",
        subject="英語",
        boxes={
            "Lesson3 ワークシート": {
                "submitted": SAMPLE_STUDENTS[:13],
                "unsubmitted": SAMPLE_STUDENTS[13:],
            },
            "Lesson4 単語テスト": {
                "submitted": SAMPLE_STUDENTS[:17],
                "unsubmitted": SAMPLE_STUDENTS[17:],
            },
        },
    )

    # --- 理科のExcel ---
    _create_subject_file(
        filepath=out / "1年A組_理科_提出状況.xlsx",
        subject="理科",
        boxes={
            "実験レポート1": {
                "submitted": SAMPLE_STUDENTS[:10],
                "unsubmitted": SAMPLE_STUDENTS[10:],
            },
        },
    )

    print(f"[INFO] サンプルExcelを生成しました: {out}/")
    for f in sorted(out.glob("*.xlsx")):
        print(f"  - {f.name}")


def _create_subject_file(filepath: Path, subject: str, boxes: dict):
    """1教科分のExcelファイルを作成する。

    ロイロノートの提出状況ダウンロード形式に似せた構造:
    - 列: 氏名, 提出箱1, 提出箱2, ...
    - 行: 各生徒
    - セル値: "提出済" or "未提出"
    """
    wb = Workbook()
    ws = wb.active
    ws.title = subject

    header_font = Font(bold=True)

    # ヘッダー行
    box_names = list(boxes.keys())
    ws.cell(row=1, column=1, value="氏名").font = header_font
    for col, name in enumerate(box_names, 2):
        ws.cell(row=1, column=col, value=name).font = header_font

    # 全生徒のリストを作成（提出済み・未提出を結合して重複除去）
    all_students = sorted(set(
        s for box in boxes.values()
        for s in box["submitted"] + box["unsubmitted"]
    ))

    # データ行
    for row, student in enumerate(all_students, 2):
        ws.cell(row=row, column=1, value=student)
        for col, box_name in enumerate(box_names, 2):
            box = boxes[box_name]
            if student in box["unsubmitted"]:
                ws.cell(row=row, column=col, value="未提出")
            else:
                ws.cell(row=row, column=col, value="提出済")

    # 列幅調整
    ws.column_dimensions["A"].width = 16
    for col in range(2, len(box_names) + 2):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = 22

    wb.save(str(filepath))


if __name__ == "__main__":
    create_sample_excel()
