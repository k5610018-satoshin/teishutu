"""未提出者レポートを生成するモジュール。

未提出者情報を教科・提出箱別に整理し、
コンソール出力またはExcelファイルとして出力します。
"""

from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def group_by_subject_and_box(unsubmitted: list[dict]) -> dict:
    """未提出者リストを 教科 → 提出箱名 → 生徒名リスト の構造に整理する。"""
    grouped = defaultdict(lambda: defaultdict(list))
    for item in unsubmitted:
        subject = item.get("subject", "不明")
        box = item.get("submission_box", "不明")
        name = item.get("student_name", "不明")
        if name not in grouped[subject][box]:
            grouped[subject][box].append(name)
    return dict(grouped)


def print_console_report(unsubmitted: list[dict]):
    """未提出者一覧をコンソールに表示する。"""
    if not unsubmitted:
        print("\n" + "=" * 50)
        print("  未提出者はいません")
        print("=" * 50)
        return

    grouped = group_by_subject_and_box(unsubmitted)
    total_count = len(unsubmitted)
    now = datetime.now().strftime("%Y/%m/%d %H:%M")

    print()
    print("=" * 60)
    print(f"  未提出者一覧  ({now} 時点)")
    print(f"  合計: {total_count} 件")
    print("=" * 60)

    for subject, boxes in sorted(grouped.items()):
        print(f"\n■ {subject}")
        print("-" * 50)
        for box_name, students in sorted(boxes.items()):
            print(f"  【{box_name}】 未提出: {len(students)}名")
            for i, name in enumerate(sorted(students), 1):
                print(f"    {i:2d}. {name}")
        print()

    print("=" * 60)


def export_excel_report(unsubmitted: list[dict], output_path: str):
    """未提出者一覧をExcelファイルに出力する。

    出力構成:
    - シート1 「未提出者一覧」: 教科・提出箱別の一覧表
    - シート2 「詳細データ」: 全データの一覧（フィルタ用）
    """
    if not unsubmitted:
        print("[INFO] 未提出者がいないため、Excelファイルは作成しません")
        return

    wb = Workbook()

    # --- シート1: 教科・提出箱別の見やすい一覧 ---
    ws1 = wb.active
    ws1.title = "未提出者一覧"
    _write_summary_sheet(ws1, unsubmitted)

    # --- シート2: 詳細データ（フィルタ用） ---
    ws2 = wb.create_sheet("詳細データ")
    _write_detail_sheet(ws2, unsubmitted)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    print(f"[INFO] Excelファイルを出力しました: {output}")


def _write_summary_sheet(ws, unsubmitted: list[dict]):
    """教科・提出箱別の一覧シートを作成する。"""
    # スタイル定義
    header_font = Font(name="メイリオ", bold=True, size=14)
    subject_font = Font(name="メイリオ", bold=True, size=12, color="FFFFFF")
    subject_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    box_font = Font(name="メイリオ", bold=True, size=11)
    box_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    normal_font = Font(name="メイリオ", size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # タイトル
    now = datetime.now().strftime("%Y/%m/%d %H:%M")
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = f"未提出者一覧 ({now} 時点)"
    title_cell.font = header_font

    row = 3
    grouped = group_by_subject_and_box(unsubmitted)

    for subject, boxes in sorted(grouped.items()):
        # 教科ヘッダー
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=1, value=f"■ {subject}")
        cell.font = subject_font
        cell.fill = subject_fill
        row += 1

        for box_name, students in sorted(boxes.items()):
            # 提出箱ヘッダー
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            cell = ws.cell(
                row=row, column=1,
                value=f"  {box_name}  (未提出: {len(students)}名)"
            )
            cell.font = box_font
            cell.fill = box_fill
            row += 1

            # 生徒名一覧（横に4列で並べる）
            cols = 4
            for i, name in enumerate(sorted(students)):
                col = (i % cols) + 1
                cell = ws.cell(row=row, column=col, value=name)
                cell.font = normal_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")
                if (i + 1) % cols == 0:
                    row += 1
            if len(students) % cols != 0:
                row += 1

            row += 1  # 提出箱間の空行

        row += 1  # 教科間の空行

    # 列幅調整
    for col_letter in ["A", "B", "C", "D"]:
        ws.column_dimensions[col_letter].width = 18


def _write_detail_sheet(ws, unsubmitted: list[dict]):
    """フィルタ・ソート用の詳細データシートを作成する。"""
    header_font = Font(name="メイリオ", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    normal_font = Font(name="メイリオ", size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # ヘッダー
    headers = ["教科", "提出箱名", "氏名"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    # データ
    for row_idx, item in enumerate(sorted(
        unsubmitted, key=lambda x: (x.get("subject", ""), x.get("submission_box", ""), x.get("student_name", ""))
    ), 2):
        for col, key in enumerate(["subject", "submission_box", "student_name"], 1):
            cell = ws.cell(row=row_idx, column=col, value=item.get(key, ""))
            cell.font = normal_font
            cell.border = thin_border

    # 列幅調整
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 16

    # オートフィルタ設定
    ws.auto_filter.ref = f"A1:C{len(unsubmitted) + 1}"
