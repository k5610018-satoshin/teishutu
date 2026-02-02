"""ロイロノートからダウンロードした提出状況Excelファイルを処理するモジュール。

ロイロノートの管理ページから「提出状況ダウンロード」で取得した
Excelファイル (submissions_status_YYYYMMDD.xlsx) を解析し、
未提出者情報を抽出します。
"""

import os
import re
from pathlib import Path

from openpyxl import load_workbook


def find_excel_files(input_dir: str) -> list[Path]:
    """指定ディレクトリ内のExcelファイルを検索する。"""
    directory = Path(input_dir)
    if not directory.exists():
        raise FileNotFoundError(f"ディレクトリが見つかりません: {input_dir}")

    xlsx_files = sorted(directory.glob("*.xlsx"))
    if not xlsx_files:
        raise FileNotFoundError(f"Excelファイルが見つかりません: {input_dir}")

    return xlsx_files


def parse_submission_excel(filepath: Path, subject: str = "") -> list[dict]:
    """ロイロノートの提出状況Excelファイルを解析する。

    ロイロノートのExcelは以下の構造を想定:
    - 1行目: ヘッダー行（提出箱名、生徒名など）
    - 各行: 生徒ごとの提出状況
    - 提出状況は「提出済」「未提出」などのテキスト

    Args:
        filepath: Excelファイルのパス
        subject: 教科名（外部から指定）

    Returns:
        未提出者情報のリスト
    """
    wb = load_workbook(filepath, read_only=True, data_only=True)
    unsubmitted = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            continue

        header = rows[0]
        if header is None:
            continue

        # ヘッダーから生徒名列と提出箱列を特定
        name_col = _find_column(header, ["名前", "氏名", "生徒名", "name"])
        if name_col is None:
            # 最初の列が生徒名の場合もある
            name_col = 0

        # 提出状況列を探す（「提出箱名」がヘッダーに含まれるパターン）
        submission_cols = _find_submission_columns(header, name_col)

        for row in rows[1:]:
            if row is None or all(cell is None for cell in row):
                continue

            student_name = str(row[name_col]).strip() if row[name_col] else ""
            if not student_name:
                continue

            for col_idx, box_name in submission_cols.items():
                if col_idx >= len(row):
                    continue
                status = str(row[col_idx]).strip() if row[col_idx] else ""
                if _is_unsubmitted(status):
                    unsubmitted.append({
                        "subject": subject or _guess_subject(filepath, sheet_name),
                        "submission_box": box_name,
                        "student_name": student_name,
                        "status": status,
                    })

    wb.close()
    return unsubmitted


def _find_column(header: tuple, keywords: list[str]) -> int | None:
    """ヘッダー行からキーワードに一致する列を探す。"""
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        cell_str = str(cell).strip()
        for kw in keywords:
            if kw in cell_str:
                return idx
    return None


def _find_submission_columns(header: tuple, name_col: int) -> dict[int, str]:
    """ヘッダー行から提出箱に対応する列を特定する。

    提出状況Excelでは、提出箱名がヘッダーに並び、
    その列に各生徒の提出状態が入るパターンが多い。
    """
    submission_cols = {}
    skip_keywords = ["名前", "氏名", "生徒名", "name", "番号", "出席", "クラス",
                     "class", "number", "page_count"]

    for idx, cell in enumerate(header):
        if idx == name_col or cell is None:
            continue
        cell_str = str(cell).strip()
        if not cell_str:
            continue
        if any(kw in cell_str.lower() for kw in skip_keywords):
            continue
        submission_cols[idx] = cell_str

    return submission_cols


def _is_unsubmitted(status: str) -> bool:
    """提出状況が「未提出」かどうか判定する。"""
    if not status or status == "None":
        return True
    unsubmitted_keywords = ["未提出", "未", "unsubmitted", "not submitted"]
    status_lower = status.lower().strip()
    for kw in unsubmitted_keywords:
        if kw in status_lower:
            return True
    # 数値の0も未提出と見なす
    if status_lower == "0":
        return True
    return False


def _guess_subject(filepath: Path, sheet_name: str) -> str:
    """ファイル名やシート名から教科名を推測する。"""
    filename = filepath.stem
    # よくある教科名パターン
    subjects = ["国語", "数学", "英語", "理科", "社会", "音楽", "美術",
                "体育", "技術", "家庭", "道徳", "総合", "情報"]
    for subj in subjects:
        if subj in filename or subj in sheet_name:
            return subj
    return filename


def process_excel_directory(input_dir: str, file_subject_map: dict = None) -> list[dict]:
    """ディレクトリ内の全Excelファイルを処理して未提出者リストを返す。

    Args:
        input_dir: Excelファイルが格納されたディレクトリ
        file_subject_map: ファイル名の一部 → 教科名 の対応辞書

    Returns:
        未提出者情報のリスト
    """
    file_subject_map = file_subject_map or {}
    files = find_excel_files(input_dir)
    all_unsubmitted = []

    for filepath in files:
        # ファイル名から教科を特定
        subject = ""
        for key, subj in file_subject_map.items():
            if key in filepath.stem:
                subject = subj
                break

        results = parse_submission_excel(filepath, subject)
        all_unsubmitted.extend(results)

    return all_unsubmitted
